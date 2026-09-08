#!/usr/bin/env python3
"""Generate report Excel dari file hasil run di results/.

Usage:
    python scripts/generate_report.py results/<modul>__<runId>.json [file2.json ...]
    python scripts/generate_report.py --all   # gabungkan run TERBARU per modul

Output: reports/<modul>__<runId>.xlsx  (atau reports/rekap-all__<timestamp>.xlsx untuk --all)
Sheet: Summary, Detail, Failed_BugCandidates
"""
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RESULTS = ROOT / "results"
REPORTS = ROOT / "reports"

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE_FONT = Font(name=FONT, size=10)
BOLD = Font(name=FONT, bold=True, size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=13)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "passed": PatternFill("solid", fgColor="C6EFCE"),
    "failed": PatternFill("solid", fgColor="FFC7CE"),
    "blocked": PatternFill("solid", fgColor="FFEB9C"),
    "skipped": PatternFill("solid", fgColor="D9D9D9"),
}

DETAIL_HEADERS = [
    "No", "Module", "Scenario ID", "Judul", "Kategori", "Prioritas", "Screen",
    "Requirements", "Status", "Durasi (s)", "Error", "Bug Candidate", "Screenshot", "Catatan",
]


def load_runs(paths):
    runs = []
    for p in paths:
        with open(p, encoding="utf-8") as f:
            runs.append(json.load(f))
    return runs


def latest_per_module():
    latest = {}
    for p in sorted(RESULTS.glob("*.json")):
        if p.name.startswith("_plan"):
            continue
        module = p.stem.split("__")[0]
        latest[module] = p  # sorted ascending -> yang terakhir menang
    return list(latest.values())


def write_detail(ws, runs):
    ws.append(DETAIL_HEADERS)
    for c in ws[1]:
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    n = 0
    for run in runs:
        for s in run.get("scenarios", []):
            n += 1
            ws.append([
                n, run.get("module", ""), s.get("id", ""), s.get("title", ""),
                s.get("category", ""), s.get("priority", ""), s.get("screen", ""),
                ", ".join(s.get("requirements", []) or []),
                s.get("status", ""), s.get("durationSec"),
                s.get("error") or "", s.get("bugCandidate") or "",
                s.get("screenshot") or "", s.get("notes") or "",
            ])
    widths = [5, 26, 12, 55, 10, 10, 22, 24, 10, 10, 45, 12, 30, 35]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BASE_FONT, BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (4, 11, 14))
        status = row[8].value  # kolom I
        if status in STATUS_FILL:
            row[8].fill = STATUS_FILL[status]
            row[8].alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{max(ws.max_row, 2)}"
    return n


def write_summary(ws, runs, total_rows):
    last = total_rows + 1  # baris data Detail terakhir
    ws["A1"] = "Laporan Hasil Automation Testing — TMS"
    ws["A1"].font = TITLE_FONT
    meta = [
        ("Digenerate", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Module / Run", "; ".join(f"{r.get('module','?')} ({r.get('runId','?')})" for r in runs)),
        ("Environment", "; ".join(sorted({(r.get("environment") or {}).get("baseUrl", "") for r in runs if r.get("environment")})) or "-"),
        ("User / Role", "; ".join(sorted({f"{(r.get('environment') or {}).get('user','')} / {(r.get('environment') or {}).get('role','')}" for r in runs if r.get("environment")})) or "-"),
    ]
    row = 3
    for label, val in meta:
        ws.cell(row=row, column=1, value=label).font = BOLD
        ws.cell(row=row, column=2, value=val).font = BASE_FONT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Status").font = HDR_FONT
    ws.cell(row=row, column=1).fill = HDR_FILL
    ws.cell(row=row, column=2, value="Jumlah").font = HDR_FONT
    ws.cell(row=row, column=2).fill = HDR_FILL
    ws.cell(row=row, column=3, value="Persentase").font = HDR_FONT
    ws.cell(row=row, column=3).fill = HDR_FILL
    start = row + 1
    for i, status in enumerate(["passed", "failed", "blocked", "skipped"]):
        r = start + i
        ws.cell(row=r, column=1, value=status.capitalize()).font = BASE_FONT
        c = ws.cell(row=r, column=2, value=f'=COUNTIF(Detail!$I$2:$I${last},"{status}")')
        c.font = BASE_FONT
        p = ws.cell(row=r, column=3, value=f"=IF($B${start+4}=0,0,B{r}/$B${start+4})")
        p.font, p.number_format = BASE_FONT, "0.0%"
    total_r = start + 4
    ws.cell(row=total_r, column=1, value="Total").font = BOLD
    ws.cell(row=total_r, column=2, value=f"=SUM(B{start}:B{start+3})").font = BOLD
    pr = ws.cell(row=total_r, column=3, value=f"=IF(B{total_r}=0,0,B{start}/B{total_r})")
    pr.font, pr.number_format = BOLD, "0.0%"
    ws.cell(row=total_r + 1, column=1, value="Catatan: persentase baris Total = pass rate.").font = Font(name=FONT, italic=True, size=9)

    # breakdown kategori x status (COUNTIFS)
    row = total_r + 3
    ws.cell(row=row, column=1, value="Breakdown per Kategori").font = BOLD
    row += 1
    headers = ["Kategori", "Passed", "Failed", "Blocked", "Skipped"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = HDR_FONT, HDR_FILL
    for i, cat in enumerate(["positive", "negative", "edge", "stress"], 1):
        r = row + i
        ws.cell(row=r, column=1, value=cat).font = BASE_FONT
        for j, status in enumerate(["passed", "failed", "blocked", "skipped"], 2):
            c = ws.cell(row=r, column=j,
                        value=f'=COUNTIFS(Detail!$E$2:$E${last},"{cat}",Detail!$I$2:$I${last},"{status}")')
            c.font = BASE_FONT
    for col, w in zip("ABCDE", [22, 40, 12, 12, 12]):
        ws.column_dimensions[col].width = w


def write_failed(ws, runs):
    headers = ["Module", "Scenario ID", "Judul", "Prioritas", "Status", "Bug Candidate (FND)",
               "Requirements", "Error / Perilaku Aktual", "Screenshot", "Catatan"]
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
    for run in runs:
        for s in run.get("scenarios", []):
            if s.get("status") == "failed" or s.get("bugCandidate"):
                ws.append([
                    run.get("module", ""), s.get("id", ""), s.get("title", ""),
                    s.get("priority", ""), s.get("status", ""), s.get("bugCandidate") or "",
                    ", ".join(s.get("requirements", []) or []),
                    s.get("error") or "", s.get("screenshot") or "", s.get("notes") or "",
                ])
    if ws.max_row == 1:
        ws.append(["-", "-", "Tidak ada failed / bug candidate pada run ini", "", "", "", "", "", "", ""])
    for i, w in enumerate([26, 12, 50, 10, 10, 16, 24, 50, 30, 35], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BASE_FONT, BORDER
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 8, 10))
    ws.freeze_panes = "A2"


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    paths = latest_per_module() if args == ["--all"] else [Path(a) for a in args]
    missing = [p for p in paths if not Path(p).exists()]
    if missing:
        sys.exit(f"File tidak ditemukan: {missing}")
    runs = load_runs(paths)
    if not runs:
        sys.exit("Tidak ada file hasil di results/.")

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_det = wb.create_sheet("Detail")
    ws_fail = wb.create_sheet("Failed_BugCandidates")

    total = write_detail(ws_det, runs)
    write_summary(ws_sum, runs, total)
    write_failed(ws_fail, runs)

    REPORTS.mkdir(exist_ok=True)
    if len(runs) == 1:
        name = f"{runs[0].get('module','run')}__{runs[0].get('runId','')}.xlsx"
    else:
        name = f"rekap-all__{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    out = REPORTS / name
    wb.save(out)
    print(f"Report tersimpan: {out}")
    print("Reminder: jalankan LibreOffice/Excel recalc bila membaca nilai formula secara programatik.")


if __name__ == "__main__":
    main()
